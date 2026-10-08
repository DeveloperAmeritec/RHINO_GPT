#!/usr/bin/env python3
"""Build OpenAI vector stores for RhinoGPT personal access profiles."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_EXCEL = "local_private/personal_matrix.xlsx"
DEFAULT_OUTPUT_DIR = "local_private/outputs/personal_vector_stores"
DEFAULT_PROFILE = "3_5"
PROFILE_PREFIX = "perfil_"
TERMINAL_BATCH_STATUSES = {"completed", "failed", "cancelled", "expired"}


def utc_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def normalize_profile_id(column_name: str) -> str:
    return column_name.removeprefix(PROFILE_PREFIX)


def truthy_cell(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().lower() in {"1", "x", "true", "si", "yes", "y"}


def parse_emails(value: Any) -> list[str]:
    if not value:
        return []
    parts = re.split(r"[,;\n]+", str(value))
    return [part.strip().lower() for part in parts if part.strip()]


def rows_from_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Missing sheet '{sheet_name}' in {path}")
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError(f"Sheet '{sheet_name}' is empty")
    clean_headers = [str(header).strip() if header is not None else "" for header in headers]
    return [
        {clean_headers[index]: value for index, value in enumerate(row)}
        for row in rows
        if any(cell is not None for cell in row)
    ]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_matrix(excel_path: Path, repo_root: Path) -> dict[str, Any]:
    matrix_rows = rows_from_sheet(excel_path, "matriz")
    profile_rows = rows_from_sheet(excel_path, "perfiles")

    profile_columns = [
        column
        for column in matrix_rows[0].keys()
        if column.startswith(PROFILE_PREFIX)
    ]
    if not profile_columns:
        raise ValueError("No profile columns found in sheet 'matriz'")

    profiles: dict[str, dict[str, Any]] = {}
    email_to_profile: dict[str, str] = {}
    duplicate_emails: dict[str, list[str]] = defaultdict(list)

    for row in profile_rows:
        profile_id = str(row.get("id_perfil") or "").strip()
        if not profile_id:
            continue
        emails = parse_emails(row.get("email"))
        profiles[profile_id] = {
            "profile_id": profile_id,
            "description": row.get("desc_perfil") or "",
            "users": row.get("usuarios") or "",
            "emails": emails,
        }
        for email in emails:
            if email in email_to_profile:
                duplicate_emails[email].extend([email_to_profile[email], profile_id])
            email_to_profile[email] = profile_id

    documents_by_profile: dict[str, list[dict[str, Any]]] = {
        normalize_profile_id(column): [] for column in profile_columns
    }
    missing_files: list[dict[str, Any]] = []
    unique_documents: dict[str, dict[str, Any]] = {}

    for row in matrix_rows:
        route = str(row.get("ruta_documento") or "").strip()
        document_name = str(row.get("nombre_documento") or "").strip()
        if not route:
            missing_files.append({"document": document_name, "reason": "empty ruta_documento"})
            continue

        absolute_path = (repo_root / route).resolve()
        document = {
            "document_type": row.get("tipo_documento"),
            "number": row.get("numero"),
            "name": document_name,
            "folder": row.get("carpeta"),
            "relative_path": route,
            "absolute_path": str(absolute_path),
            "match_status": row.get("estado_match"),
            "match_notes": row.get("notas_match"),
        }

        if not absolute_path.exists():
            missing_files.append({**document, "reason": "file not found"})
        else:
            document["sha256"] = sha256_file(absolute_path)
            unique_documents[str(absolute_path)] = document

        for column in profile_columns:
            if truthy_cell(row.get(column)):
                profile_id = normalize_profile_id(column)
                documents_by_profile[profile_id].append(document)

    for profile_id in documents_by_profile:
        profiles.setdefault(
            profile_id,
            {
                "profile_id": profile_id,
                "description": "",
                "users": "",
                "emails": [],
            },
        )

    return {
        "profiles": profiles,
        "email_to_profile": email_to_profile,
        "duplicate_emails": {
            email: sorted(set(profile_ids))
            for email, profile_ids in duplicate_emails.items()
        },
        "documents_by_profile": documents_by_profile,
        "unique_documents": unique_documents,
        "missing_files": missing_files,
    }


def get_openai_client():
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("API_OPENAI")
    if not api_key:
        raise RuntimeError("Set OPENAI_API_KEY or API_OPENAI before using --execute")
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("Install the OpenAI Python SDK before using --execute") from exc
    return OpenAI(api_key=api_key)


def get_vector_stores_resource(client):
    if hasattr(client, "vector_stores"):
        return client.vector_stores
    beta = getattr(client, "beta", None)
    if beta and hasattr(beta, "vector_stores"):
        return beta.vector_stores
    raise RuntimeError("OpenAI SDK does not expose vector_stores")


def create_file(client, path: Path):
    with path.open("rb") as handle:
        return client.files.create(file=handle, purpose="assistants")


def wait_for_batch(vector_stores, vector_store_id: str, batch_id: str, timeout_seconds: int):
    deadline = time.time() + timeout_seconds
    while True:
        batch = vector_stores.file_batches.retrieve(
            vector_store_id=vector_store_id,
            batch_id=batch_id,
        )
        if batch.status in TERMINAL_BATCH_STATUSES:
            return batch
        if time.time() > deadline:
            raise TimeoutError(
                f"Timed out waiting for batch {batch_id} on vector store {vector_store_id}"
            )
        time.sleep(5)


def chunked(values: list[str], size: int) -> list[list[str]]:
    return [values[index : index + size] for index in range(0, len(values), size)]


def list_existing_vector_stores_by_name(vector_stores, names: set[str]) -> dict[str, list[str]]:
    matches: dict[str, list[str]] = defaultdict(list)
    page = vector_stores.list(limit=100)
    while True:
        for item in getattr(page, "data", []):
            if item.name in names:
                matches[item.name].append(item.id)
        data = getattr(page, "data", [])
        if not getattr(page, "has_more", False) or not data:
            break
        page = vector_stores.list(limit=100, after=data[-1].id)
    return dict(matches)


def execute_build(
    matrix: dict[str, Any],
    repo_root: Path,
    vector_store_prefix: str,
    batch_timeout_seconds: int,
    include_timestamp: bool,
    replace_existing_stores: bool,
) -> dict[str, Any]:
    client = get_openai_client()
    vector_stores = get_vector_stores_resource(client)

    planned_names = {
        f"{vector_store_prefix}_{profile_id}"
        for profile_id in matrix["documents_by_profile"]
    }
    if include_timestamp:
        planned_names = set()

    existing_stores = (
        list_existing_vector_stores_by_name(vector_stores, planned_names)
        if planned_names
        else {}
    )
    if existing_stores and not replace_existing_stores:
        names = ", ".join(sorted(existing_stores))
        raise RuntimeError(
            f"Vector store name(s) already exist: {names}. "
            "Use --replace-existing-stores to delete and recreate them."
        )
    for name, vector_store_ids in sorted(existing_stores.items()):
        for vector_store_id in vector_store_ids:
            vector_stores.delete(vector_store_id=vector_store_id)
            print(f"deleted existing {name} -> {vector_store_id}")

    file_ids_by_path: dict[str, str] = {}
    uploaded_files: dict[str, dict[str, Any]] = {}

    for absolute_path, document in sorted(matrix["unique_documents"].items()):
        created_file = create_file(client, Path(absolute_path))
        file_ids_by_path[absolute_path] = created_file.id
        uploaded_files[absolute_path] = {
            "file_id": created_file.id,
            "relative_path": document["relative_path"],
            "name": document["name"],
            "sha256": document.get("sha256"),
        }
        print(f"uploaded {document['relative_path']} -> {created_file.id}")

    vector_store_results: dict[str, dict[str, Any]] = {}
    stamp = utc_stamp()

    for profile_id, documents in sorted(matrix["documents_by_profile"].items()):
        file_ids = sorted(
            {
                file_ids_by_path[document["absolute_path"]]
                for document in documents
                if document["absolute_path"] in file_ids_by_path
            }
        )
        vector_store_name = f"{vector_store_prefix}_{profile_id}"
        if include_timestamp:
            vector_store_name = f"{vector_store_name}_{stamp}"

        vector_store = vector_stores.create(
            name=vector_store_name,
            metadata={
                "source": "local_private/personal_matrix.xlsx",
                "profile_id": profile_id,
                "default_profile": str(profile_id == DEFAULT_PROFILE).lower(),
                "repo": repo_root.name,
            },
        )

        batches: list[dict[str, Any]] = []
        for file_id_chunk in chunked(file_ids, 500):
            batch = vector_stores.file_batches.create(
                vector_store_id=vector_store.id,
                file_ids=file_id_chunk,
            )
            batch = wait_for_batch(
                vector_stores,
                vector_store.id,
                batch.id,
                batch_timeout_seconds,
            )
            batches.append(
                {
                    "batch_id": batch.id,
                    "status": batch.status,
                    "file_counts": getattr(batch, "file_counts", None).model_dump()
                    if getattr(batch, "file_counts", None)
                    else None,
                }
            )

        vector_store_results[profile_id] = {
            "vector_store_id": vector_store.id,
            "name": vector_store.name,
            "document_count": len(documents),
            "file_count": len(file_ids),
            "file_ids": file_ids,
            "batches": batches,
            "documents": [
                {
                    "document_type": document["document_type"],
                    "number": document["number"],
                    "name": document["name"],
                    "relative_path": document["relative_path"],
                    "file_id": file_ids_by_path.get(document["absolute_path"]),
                }
                for document in documents
            ],
        }
        print(f"created {vector_store.name} -> {vector_store.id} ({len(file_ids)} files)")

    return {
        "uploaded_files": uploaded_files,
        "vector_stores": vector_store_results,
    }


def build_output_payload(
    matrix: dict[str, Any],
    excel_path: Path,
    repo_root: Path,
    execute_result: dict[str, Any] | None,
) -> dict[str, Any]:
    profiles = matrix["profiles"]
    documents_by_profile = matrix["documents_by_profile"]

    payload = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "source_excel": str(excel_path),
        "default_profile": DEFAULT_PROFILE,
        "profile_count": len(profiles),
        "unique_document_count": len(matrix["unique_documents"]),
        "missing_file_count": len(matrix["missing_files"]),
        "profiles": profiles,
        "email_to_profile": matrix["email_to_profile"],
        "duplicate_emails": matrix["duplicate_emails"],
        "profile_summary": {
            profile_id: {
                "document_count": len(documents),
                "description": profiles.get(profile_id, {}).get("description", ""),
                "email_count": len(profiles.get(profile_id, {}).get("emails", [])),
            }
            for profile_id, documents in sorted(documents_by_profile.items())
        },
        "missing_files": matrix["missing_files"],
        "documents_by_profile": {
            profile_id: [
                {
                    "document_type": document["document_type"],
                    "number": document["number"],
                    "name": document["name"],
                    "relative_path": document["relative_path"],
                    "match_status": document["match_status"],
                    "sha256": document.get("sha256"),
                }
                for document in documents
            ]
            for profile_id, documents in sorted(documents_by_profile.items())
        },
    }

    if execute_result:
        payload.update(execute_result)

    return payload


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create one OpenAI vector store per personal profile from the private local matrix."
    )
    parser.add_argument("--excel", default=DEFAULT_EXCEL, help="Path to the matrix workbook.")
    parser.add_argument(
        "--output-dir",
        default=DEFAULT_OUTPUT_DIR,
        help="Directory for the JSON result.",
    )
    parser.add_argument(
        "--execute",
        action="store_true",
        help="Actually upload files and create vector stores. Default is dry-run.",
    )
    parser.add_argument(
        "--allow-missing-files",
        action="store_true",
        help="Continue even if some paths from the workbook do not exist.",
    )
    parser.add_argument(
        "--vector-store-prefix",
        default="rhinogpt_personal_profile",
        help="Prefix for created OpenAI vector store names.",
    )
    parser.add_argument(
        "--include-timestamp",
        action="store_true",
        help="Append a UTC timestamp to vector store names. Off by default for stable production names.",
    )
    parser.add_argument(
        "--replace-existing-stores",
        action="store_true",
        help="Delete existing vector stores with the planned stable names before creating new ones.",
    )
    parser.add_argument(
        "--batch-timeout-seconds",
        type=int,
        default=1800,
        help="Maximum seconds to wait for each OpenAI file batch.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    repo_root = Path.cwd().resolve()
    excel_path = (repo_root / args.excel).resolve()
    if not excel_path.exists():
        print(f"Excel file not found: {excel_path}", file=sys.stderr)
        return 1

    matrix = build_matrix(excel_path, repo_root)
    if matrix["missing_files"] and not args.allow_missing_files:
        print("Missing files found. Re-run with --allow-missing-files only if intentional.", file=sys.stderr)
        for item in matrix["missing_files"][:20]:
            print(f"- {item.get('relative_path') or item.get('document')}: {item['reason']}", file=sys.stderr)
        if len(matrix["missing_files"]) > 20:
            print(f"... and {len(matrix['missing_files']) - 20} more", file=sys.stderr)
        return 1

    execute_result = None
    if args.execute:
        execute_result = execute_build(
            matrix,
            repo_root,
            args.vector_store_prefix,
            args.batch_timeout_seconds,
            args.include_timestamp,
            args.replace_existing_stores,
        )

    output_dir = (repo_root / args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"personal_vector_stores_{utc_stamp()}.json"
    payload = build_output_payload(matrix, excel_path, repo_root, execute_result)
    payload["mode"] = "execute" if args.execute else "dry_run"
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps(payload["profile_summary"], indent=2, ensure_ascii=False))
    print(f"wrote {output_path}")
    if not args.execute:
        print("dry-run only; add --execute after setting OPENAI_API_KEY or API_OPENAI")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
