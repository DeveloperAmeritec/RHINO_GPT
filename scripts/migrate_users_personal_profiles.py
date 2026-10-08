#!/usr/bin/env python3
"""Migrate RhinoGPT Users to the lean schema and load personal profiles."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

import boto3
from boto3.dynamodb.types import TypeDeserializer
from openpyxl import load_workbook


DEFAULT_EXCEL = "local_private/personal_matrix.xlsx"
DEFAULT_OUTPUT_DIR = "local_private/outputs/users_migration"
DEFAULT_MONTHLY_LIMIT = 100
DEFAULT_PERSONAL_PROFILE_ID = "3_5"
LEGACY_FIELDS = ["is_active", "organization", "plan"]


def get_gt_time() -> str:
    utc_time = datetime.now(timezone.utc)
    gt_time = utc_time - timedelta(hours=6)
    return gt_time.strftime("%Y-%m-%dT%H:%M:%S-06:00")


def parse_list(value: Any) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in re.split(r"[,;\n]+", str(value)) if part.strip()]


def parse_names(value: Any, expected_count: int) -> list[str]:
    names = parse_list(value)
    while len(names) < expected_count:
        split_index = next(
            (
                index
                for index, name in enumerate(names)
                if " y " in name
            ),
            None,
        )
        if split_index is None:
            break
        left, right = names[split_index].rsplit(" y ", 1)
        names[split_index : split_index + 1] = [left.strip(), right.strip()]
    return names


def load_personal_users(excel_path: Path) -> list[dict[str, str]]:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if "perfiles" not in wb.sheetnames:
        raise ValueError("Missing sheet 'perfiles'")
    ws = wb["perfiles"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]

    users: list[dict[str, str]] = []
    for row in rows:
        item = {headers[index]: value for index, value in enumerate(row)}
        profile_id = str(item.get("id_perfil") or "").strip()
        emails = [email.lower() for email in parse_list(item.get("email"))]
        names = parse_names(item.get("usuarios"), len(emails))

        for index, email in enumerate(emails):
            users.append(
                {
                    "user_id": email,
                    "full_name": names[index] if index < len(names) else email,
                    "user_type": "personal",
                    "personal_profile_id": profile_id or DEFAULT_PERSONAL_PROFILE_ID,
                }
            )
    return users


def scan_table(client, table_name: str) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    params = {"TableName": table_name}
    while True:
        response = client.scan(**params)
        items.extend(response.get("Items", []))
        if "LastEvaluatedKey" not in response:
            break
        params["ExclusiveStartKey"] = response["LastEvaluatedKey"]
    return items


def save_backup(items: list[dict[str, Any]], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"users_backup_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.json"
    deserializer = TypeDeserializer()
    plain_items = [
        {key: deserializer.deserialize(value) for key, value in item.items()}
        for item in items
    ]
    output_path.write_text(
        json.dumps(plain_items, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    return output_path


def remove_legacy_fields(client, table_name: str, items: list[dict[str, Any]], execute: bool) -> int:
    updated = 0
    for item in items:
        removable = [field for field in LEGACY_FIELDS if field in item]
        if not removable:
            continue
        updated += 1
        if execute:
            attribute_names = {
                f"#field{index}": field for index, field in enumerate(removable)
            }
            client.update_item(
                TableName=table_name,
                Key={"user_id": item["user_id"]},
                UpdateExpression="REMOVE " + ", ".join(attribute_names),
                ExpressionAttributeNames=attribute_names,
            )
    return updated


def upsert_personal_users(
    client,
    table_name: str,
    users: list[dict[str, str]],
    monthly_limit: int,
    execute: bool,
) -> int:
    now = get_gt_time()
    for user in users:
        if execute:
            client.update_item(
                TableName=table_name,
                Key={"user_id": {"S": user["user_id"]}},
                UpdateExpression=(
                    "SET full_name = :full_name, "
                    "user_type = :user_type, "
                    "personal_profile_id = :profile_id, "
                    "monthly_limit = if_not_exists(monthly_limit, :monthly_limit), "
                    "current_usage = if_not_exists(current_usage, :zero), "
                    "reset_date = if_not_exists(reset_date, :now), "
                    "registration_date = if_not_exists(registration_date, :now), "
                    "last_access_date = if_not_exists(last_access_date, :now) "
                    "REMOVE is_active, organization, #plan"
                ),
                ExpressionAttributeNames={"#plan": "plan"},
                ExpressionAttributeValues={
                    ":full_name": {"S": user["full_name"]},
                    ":user_type": {"S": "personal"},
                    ":profile_id": {"S": user["personal_profile_id"]},
                    ":monthly_limit": {"N": str(monthly_limit)},
                    ":zero": {"N": "0"},
                    ":now": {"S": now},
                },
            )
    return len(users)


def set_default_profile_for_existing_personal(
    client,
    table_name: str,
    items: list[dict[str, Any]],
    execute: bool,
) -> int:
    updated = 0
    for item in items:
        if item.get("user_type", {}).get("S") != "personal":
            continue
        if "personal_profile_id" in item:
            continue
        updated += 1
        if execute:
            client.update_item(
                TableName=table_name,
                Key={"user_id": item["user_id"]},
                UpdateExpression="SET personal_profile_id = :profile_id",
                ExpressionAttributeValues={
                    ":profile_id": {"S": DEFAULT_PERSONAL_PROFILE_ID},
                },
            )
    return updated


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Clean Users schema and load personal profiles.")
    parser.add_argument("--excel", default=DEFAULT_EXCEL)
    parser.add_argument("--table", default="Users")
    parser.add_argument("--region", default="us-east-2")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--monthly-limit", type=int, default=DEFAULT_MONTHLY_LIMIT)
    parser.add_argument("--execute", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    excel_path = Path(args.excel).resolve()
    if not excel_path.exists():
        raise FileNotFoundError(excel_path)

    client = boto3.client("dynamodb", region_name=args.region)
    current_items = scan_table(client, args.table)
    backup_path = save_backup(current_items, Path(args.output_dir).resolve())
    personal_users = load_personal_users(excel_path)

    remove_count = remove_legacy_fields(client, args.table, current_items, args.execute)
    default_count = set_default_profile_for_existing_personal(
        client,
        args.table,
        current_items,
        args.execute,
    )
    upsert_count = upsert_personal_users(
        client,
        args.table,
        personal_users,
        args.monthly_limit,
        args.execute,
    )

    summary = {
        "mode": "execute" if args.execute else "dry_run",
        "backup_path": str(backup_path),
        "items_scanned": len(current_items),
        "legacy_field_items_to_update": remove_count,
        "existing_personal_default_profiles_to_set": default_count,
        "personal_users_from_excel": upsert_count,
        "personal_profiles": sorted({user["personal_profile_id"] for user in personal_users}),
    }
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
