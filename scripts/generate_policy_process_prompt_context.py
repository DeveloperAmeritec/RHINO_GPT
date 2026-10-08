#!/usr/bin/env python3
"""Generate a compact policy-process context block for the personal prompt."""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_EXCEL = "local_private/knowledge/personal/Index_SmartSchool_Policies and Processes.xlsx"
DEFAULT_OUTPUT = "app/prompts/policy_process_map_context.txt"
DEFAULT_SHEET = "General"


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).strip()


def split_policies(value: str) -> list[str]:
    return [policy.strip() for policy in value.split("/") if policy.strip()]


def build_context(excel_path: Path, sheet_name: str) -> str:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {excel_path}")

    worksheet = workbook[sheet_name]
    policy_to_processes: dict[str, list[str]] = defaultdict(list)
    processes_without_policy: list[str] = []

    for row in worksheet.iter_rows(min_row=3, values_only=True):
        process = clean(row[1] if len(row) > 1 else "")
        policies = clean(row[3] if len(row) > 3 else "")
        if not process:
            continue
        if not policies:
            processes_without_policy.append(process)
            continue

        for policy in split_policies(policies):
            if process not in policy_to_processes[policy]:
                policy_to_processes[policy].append(process)

    lines = [
        "Mapa interno de relacion politica-proceso.",
        "Usa este mapa solo para identificar que politica se relaciona con un proceso o que procesos dependen de una politica.",
        "Si el usuario pide contenido, pasos o detalles de la politica/proceso, usa file_search y respeta los vector stores disponibles para su perfil.",
        "",
        "Politicas y procesos relacionados:",
    ]

    for policy in sorted(policy_to_processes):
        processes = "; ".join(sorted(policy_to_processes[policy]))
        lines.append(f"- {policy}: {processes}.")

    if processes_without_policy:
        lines.extend(
            [
                "",
                "Procesos sin politica asociada en la matriz General:",
            ]
        )
        for process in sorted(processes_without_policy):
            lines.append(f"- {process}.")

    return "\n".join(lines).strip() + "\n"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create prompt context from the General policy-process matrix."
    )
    parser.add_argument("--excel", default=DEFAULT_EXCEL)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    excel_path = Path(args.excel).resolve()
    output_path = Path(args.output).resolve()
    context = build_context(excel_path, args.sheet)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(context, encoding="utf-8")
    print(f"Wrote {output_path}")
    print(f"Lines: {len(context.splitlines())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
